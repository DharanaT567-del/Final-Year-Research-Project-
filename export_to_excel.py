# -*- coding: utf-8 -*-
import csv
import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Paths
PATH_STRICT_CSV = r'c:\Users\thiyu\Desktop\Final-Year-Research-Project-\mathaka_classical_maestros.csv'
PATH_GOLDEN_CSV = r'c:\Users\thiyu\Desktop\Final-Year-Research-Project-\mathaka_classical_and_golden_era.csv'

OUTPUT_EXCEL_COMBINED = r'c:\Users\thiyu\Desktop\Final-Year-Research-Project-\mathaka_sinhala_classical_songs.xlsx'
OUTPUT_EXCEL_STRICT = r'c:\Users\thiyu\Desktop\Final-Year-Research-Project-\mathaka_classical_maestros.xlsx'

def read_csv(file_path):
    with open(file_path, 'r', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        return reader.fieldnames, list(reader)

def style_header_row(ws, col_count):
    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='1F4E79'),
        bottom=Side(style='medium', color='1F4E79')
    )
    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

def create_summary_sheet(wb, strict_rows, golden_rows):
    ws = wb.create_sheet(title='Summary & Artist Stats', index=0)
    
    # Title
    ws.cell(row=1, column=1, value='Sinhala Classical Songs Dataset Summary').font = Font(name='Calibri', size=16, bold=True, color='1F4E79')
    
    # General Stats
    ws.cell(row=3, column=1, value='Metric').font = Font(bold=True)
    ws.cell(row=3, column=2, value='Count').font = Font(bold=True)
    
    ws.cell(row=4, column=1, value='Strict Classical Maestros Songs')
    ws.cell(row=4, column=2, value=len(strict_rows))
    
    ws.cell(row=5, column=1, value='Classical + Golden Era Melodic Songs')
    ws.cell(row=5, column=2, value=len(golden_rows))

    # Artist Frequency Table for Strict Classical
    ws.cell(row=7, column=1, value='Artist Name').font = Font(bold=True)
    ws.cell(row=7, column=2, value='Song Count (Strict Classical)').font = Font(bold=True)

    counts = {}
    for r in strict_rows:
        art = r.get('artist_name', '').strip()
        counts[art] = counts.get(art, 0) + 1
    
    row_idx = 8
    for art, c in sorted(counts.items(), key=lambda x: x[1], reverse=True):
        ws.cell(row=row_idx, column=1, value=art)
        ws.cell(row=row_idx, column=2, value=c)
        row_idx += 1

    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 30

def populate_sheet(ws, fieldnames, rows):
    ws.append(fieldnames)
    style_header_row(ws, len(fieldnames))

    for r_idx, row in enumerate(rows, start=2):
        row_vals = [row.get(fn, '') for fn in fieldnames]
        ws.append(row_vals)
        ws.row_dimensions[r_idx].height = 20

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

def main():
    fields_s, strict_rows = read_csv(PATH_STRICT_CSV)
    fields_g, golden_rows = read_csv(PATH_GOLDEN_CSV)

    # 1. Combined Excel Workbook
    wb_combined = openpyxl.Workbook()
    ws_strict = wb_combined.active
    ws_strict.title = 'Strict Classical Maestros'
    populate_sheet(ws_strict, fields_s, strict_rows)

    ws_golden = wb_combined.create_sheet(title='Classical & Golden Era')
    populate_sheet(ws_golden, fields_g, golden_rows)

    create_summary_sheet(wb_combined, strict_rows, golden_rows)

    wb_combined.save(OUTPUT_EXCEL_COMBINED)
    print(f'Combined Excel created: {OUTPUT_EXCEL_COMBINED}')

    # 2. Strict Classical Only Excel Workbook
    wb_strict = openpyxl.Workbook()
    ws = wb_strict.active
    ws.title = 'Classical Maestros'
    populate_sheet(ws, fields_s, strict_rows)
    wb_strict.save(OUTPUT_EXCEL_STRICT)
    print(f'Strict Classical Excel created: {OUTPUT_EXCEL_STRICT}')

if __name__ == '__main__':
    main()
